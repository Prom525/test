#!/usr/bin/env python3
"""
ingest_totaal_to_calendar_v0.py

Leest "TOTAAL.xlsx" (jaar-tabs) en schrijft per week-cel records naar:
  sb_maintenance_calendar_v0

Interpretatie:
- Sheetnaam = jaar (bv "2019") als parsebaar
- Weeks staan op rij 3 vanaf kolom E (E=5)
- Vanaf rij 4 naar beneden:
    A: hoofdlijn (line_name)  (kan leeg zijn; dan "carry forward")
    B: transportband (conveyor) (kan leeg zijn; carry forward)
    C: bandbreedte (belt_width) (kan leeg zijn; carry forward)
    D: schraper (scraper) (kan leeg zijn; carry forward)
  Vanaf kolom E: cellen met X = performed
  Als cel rood (fill) = blocked

Run:
  $env:DB_URL = "postgresql://postgres:SterkWachtwoord123@127.0.0.1:15432/promati"
  python .\ingest_totaal_to_calendar_v0.py "C:\ai-platform\data\TATA\Lijsten GSL\TOTAAL.xlsx" --reset
"""

import argparse
import os
import re
from typing import Optional

from openpyxl import load_workbook
from sqlalchemy import create_engine, text

DB_URL = os.environ.get("DB_URL", "postgresql://postgres:SterkWachtwoord123@127.0.0.1:15432/promati")

RE_YEAR = re.compile(r"^\s*(19|20)\d{2}\s*$")


def excel_fill_rgb(cell) -> Optional[str]:
    """Return ARGB like 'FFFF0000' if available."""
    try:
        fill = cell.fill
        if not fill or not getattr(fill, "patternType", None):
            return None
        fg = getattr(fill, "fgColor", None)
        if not fg:
            return None
        rgb = getattr(fg, "rgb", None)
        if rgb:
            return str(rgb).upper()
        return None
    except Exception:
        return None


def is_red(rgb: Optional[str]) -> bool:
    """Heuristiek voor rood."""
    if not rgb:
        return False
    rgb = rgb.upper()
    if rgb in ("FFFF0000", "FF0000", "00FF0000"):
        return True
    hex6 = rgb[-6:]
    if not re.fullmatch(r"[0-9A-F]{6}", hex6):
        return False
    r = int(hex6[0:2], 16)
    g = int(hex6[2:4], 16)
    b = int(hex6[4:6], 16)
    return r >= 200 and g <= 80 and b <= 80


def normalize_text(v) -> str:
    if v is None:
        return ""
    s = str(v).strip()
    if s.endswith(".0"):
        try:
            f = float(s)
            if f.is_integer():
                s = str(int(f))
        except Exception:
            pass
    return s


def ensure_table(conn):
    conn.execute(text("""
    create table if not exists sb_maintenance_calendar_v0 (
      id bigserial primary key,

      source_file text not null,
      sheet text not null,
      year int,

      line_name text,
      conveyor text,
      belt_width text,
      scraper text,
      row_nr int not null,

      week_nr int not null,
      col_nr int not null,

      performed boolean not null default false,
      blocked boolean not null default false,
      cell_text text,
      fill_rgb text,

      created_at timestamptz default now(),

      unique (source_file, sheet, row_nr, week_nr)
    );
    """))
    conn.execute(text("create index if not exists ix_sb_cal_year_week on sb_maintenance_calendar_v0(year, week_nr)"))
    conn.execute(text("create index if not exists ix_sb_cal_source on sb_maintenance_calendar_v0(source_file)"))
    conn.execute(text("create index if not exists ix_sb_cal_line on sb_maintenance_calendar_v0(line_name)"))


def parse_week_header(ws, start_col=5, header_row=3) -> dict[int, int]:
    """Return mapping col_nr -> week_nr from row 3, starting at column E."""
    col_to_week: dict[int, int] = {}
    max_col = ws.max_column
    for col in range(start_col, max_col + 1):
        v = ws.cell(row=header_row, column=col).value
        s = normalize_text(v)
        if not s:
            continue
        m = re.search(r"(\d{1,2})", s)
        if not m:
            continue
        wk = int(m.group(1))
        if 1 <= wk <= 53:
            col_to_week[col] = wk
    return col_to_week


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("xlsx", help="Pad naar TOTAAL.xlsx")
    ap.add_argument("--reset", action="store_true", help="truncate tabel voor herladen")
    ap.add_argument("--only-year", type=int, default=None, help="Alleen deze sheet/jaar laden (bv 2019)")
    args = ap.parse_args()

    xlsx_path = args.xlsx
    if not os.path.exists(xlsx_path):
        raise SystemExit(f"Bestand niet gevonden: {xlsx_path}")

    engine = create_engine(DB_URL)

    with engine.begin() as conn:
        ensure_table(conn)
        if args.reset:
            conn.execute(text("truncate table sb_maintenance_calendar_v0"))
            print("TRUNCATED sb_maintenance_calendar_v0")

    wb = load_workbook(xlsx_path, data_only=True)

    upsert_sql = text("""
      insert into sb_maintenance_calendar_v0
        (source_file, sheet, year, line_name, conveyor, belt_width, scraper,
         row_nr, week_nr, col_nr, performed, blocked, cell_text, fill_rgb)
      values
        (:source_file, :sheet, :year, :line_name, :conveyor, :belt_width, :scraper,
         :row_nr, :week_nr, :col_nr, :performed, :blocked, :cell_text, :fill_rgb)
      on conflict (source_file, sheet, row_nr, week_nr) do update
        set line_name = excluded.line_name,
            conveyor = excluded.conveyor,
            belt_width = excluded.belt_width,
            scraper = excluded.scraper,
            col_nr = excluded.col_nr,
            performed = excluded.performed,
            blocked = excluded.blocked,
            cell_text = excluded.cell_text,
            fill_rgb = excluded.fill_rgb
    """)

    inserted = 0

    for sheet_name in wb.sheetnames:
        if args.only_year is not None and str(args.only_year).strip() != str(sheet_name).strip():
            continue

        ws = wb[sheet_name]
        year = int(str(sheet_name).strip()) if RE_YEAR.match(str(sheet_name)) else None

        col_to_week = parse_week_header(ws, start_col=5, header_row=3)
        if not col_to_week:
            print(f"SKIP sheet '{sheet_name}': geen weekkolommen gevonden op rij 3 vanaf kolom E")
            continue

        cur_line = ""
        cur_conv = ""
        cur_width = ""
        cur_scraper = ""

        for r in range(4, ws.max_row + 1):
            a = normalize_text(ws.cell(row=r, column=1).value)
            b = normalize_text(ws.cell(row=r, column=2).value)
            c = normalize_text(ws.cell(row=r, column=3).value)
            d = normalize_text(ws.cell(row=r, column=4).value)

            if a:
                cur_line = a
            if b:
                cur_conv = b
            if c:
                cur_width = c
            if d:
                cur_scraper = d

            if not (cur_line or cur_conv or cur_width or cur_scraper):
                continue

            for col_nr, week_nr in col_to_week.items():
                cell = ws.cell(row=r, column=col_nr)
                raw = normalize_text(cell.value)
                performed = raw.upper() == "X"
                rgb = excel_fill_rgb(cell)
                blocked = is_red(rgb)

                if not performed and not blocked and not raw:
                    continue

                with engine.begin() as conn:
                    conn.execute(
                        upsert_sql,
                        {
                            "source_file": xlsx_path,
                            "sheet": sheet_name,
                            "year": year,
                            "line_name": cur_line or None,
                            "conveyor": cur_conv or None,
                            "belt_width": cur_width or None,
                            "scraper": cur_scraper or None,
                            "row_nr": r,
                            "week_nr": week_nr,
                            "col_nr": col_nr,
                            "performed": performed,
                            "blocked": blocked,
                            "cell_text": raw or None,
                            "fill_rgb": rgb,
                        },
                    )
                inserted += 1

        print(f"DONE sheet '{sheet_name}' (year={year})")

    print(f"DONE. records upserted: {inserted}")


if __name__ == "__main__":
    main()
